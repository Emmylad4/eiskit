"""Create example_materials.xlsx: five synthetic catalysts measured at one potential, side by side
(the layout of a typical comparison figure). Realistic noise, a few glitches and an inductive
high-frequency tail are included – the robust fit handles them without deleting anything."""
import numpy as np
import openpyxl
import eiskit as ek


def make(path="example_materials.xlsx", seed=0):
    rng = np.random.default_rng(seed)
    f = np.logspace(5, -1, 61)
    c = ek.Circuit("L0-R0-p(R1,CPE1)")
    samples = {"NF/Ni(OH)2": [12.5, 2e-3], "NF/Ni3N": [8.2, 3e-3], "NF/Ni3N@PA": [4.6, 5e-3],
               "NF/Ni3N@NiFe-PA": [1.0, 2e-2], "NF/NiFe-PA": [2.1, 1e-2]}
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "EIS"
    for k, (name, (Rct, Q)) in enumerate(samples.items()):
        Z = c([3e-7, 1.0 + 0.05 * k, Rct, Q, 0.85], f)
        Z = Z + 0.004 * abs(Z) * (rng.normal(size=f.size) + 1j * rng.normal(size=f.size))
        Z[[22, 37]] *= np.array([1.05, 0.93])                     # two glitches per sample
        col = 1 + 4 * k
        ws.cell(1, col, name)
        for j, h in enumerate(["Frequency (Hz)", "Z' (Ohm)", "-Z'' (Ohm)"]):
            ws.cell(2, col + j, h)
        for i, (fi, zi) in enumerate(zip(f, Z)):
            for j, v in enumerate((fi, zi.real, -zi.imag)):
                ws.cell(3 + i, col + j, float(v))
    wb.save(path)
    return path


if __name__ == "__main__":
    print(make())
