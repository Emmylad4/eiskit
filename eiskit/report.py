"""
Publication tables for :class:`~eiskit.fitting.Results` (several samples).

* :func:`tables`           – Table 1 (elements & values per sample), S1 circuits, S2 goodness
  of fit + Kramers–Kronig, S3 95 % confidence intervals, S4 derived quantities, outliers
* :func:`comparison_table` – elements as rows and samples as columns (value ± SD)
* :func:`to_excel`         – one workbook with all tables (confidence intervals as live formulas)
* :func:`to_docx`          – Word report with formatted tables (needs ``python-docx``)

Text markup in the tables: ``_x`` = subscript, ``^x`` = superscript, ``^a`` = footnote
"not determined" (relative standard error > 100 %).
"""
import math
import re

import numpy as np
import pandas as pd

UNIT = {"Ohm": "Ω", "F": "F", "H": "H", "S*s^n": "S·s^n", "s": "s", "Ohm*s^-0.5": "Ω·s^-0.5", "-": ""}
_PART = {"Q": "-T", "n": "-P", "R": "-R", "tau": "-T", "sigma": "-σ", "Rion": "-R_ion", "Rct": "-R_ct"}


def element_label(name, labels=None):
    """Table label of a parameter: R0/Rs -> R_s, R1 -> R_1, Rct -> R_ct, CPE1_Q -> CPE1-T,
    CPE1_n -> CPE1-P, Ws1_tau -> Ws1-T. ``labels`` overrides element names, e.g. {"R2": "R_ct"}."""
    labels = labels or {}
    el, _, part = name.partition("_")
    if el in labels:
        base = labels[el]
    elif el in ("R0", "Rs"):
        base = "R_s"
    else:
        m = re.fullmatch(r"([RCL])(\d+|ct|dl|sei|film|pore|ion|ads|sol|ox|po|gb|ad|el|b|f)", el)
        base = f"{m.group(1)}_{m.group(2)}" if m else el
    return base + _PART.get(part, "")


def sci(x, n=2):
    if x == 0 or not np.isfinite(x):
        return f"{x:g}"
    e = int(math.floor(math.log10(abs(x))))
    if -2 <= e <= 3:
        return f"{x:.{n}g}"
    m, ee = f"{x:.{n - 1}e}".split("e")
    return f"{m} × 10^{int(ee)}"


def value_pm(v, s):
    """'value ± SD' with two significant figures on the SD."""
    if not np.isfinite(s) or s <= 1e-6 * abs(v):
        return sci(v, 2)
    d = max(0, -int(math.floor(math.log10(abs(s)))) + 1)
    if abs(v) < 1e-3 or abs(v) >= 1e4:
        e = int(math.floor(math.log10(abs(v))))
        m, ms = v / 10 ** e, s / 10 ** e
        dm = max(0, -int(math.floor(math.log10(ms))) + 1)
        return f"({m:.{dm}f} ± {ms:.{dm}f}) × 10^{e}"
    return f"{v:.{d}f} ± {s:.{d}f}"


def tables(results, labels=None):
    """Dictionary of DataFrames: table1, comparison, S1, S2, S3, S4, outliers."""
    t1, s2, s3, s4, out = [], [], [], [], []
    for name, r in results.fits.items():
        kk = results.kk.get(name)
        for nm, v, s, rel, u in zip(r.circuit.names, r.p, r.sd, r.rel, r.circuit.units):
            lab, un = element_label(nm, labels), UNIT.get(u, u)
            val = value_pm(v, s) + (f" {un}" if un else "")
            if not np.isfinite(rel) or rel > 100:
                val += " ^a"
            t1.append({"Catalyst": name, "Element": lab, "Value": val})
            s3.append({"Sample": name, "Element": lab, "Unit": un, "Value": v, "SD": s, "Rel. error (%)": rel,
                       "95 % CI low": v - 1.96 * s, "95 % CI high": v + 1.96 * s,
                       "Determined": bool(np.isfinite(rel) and rel <= 100)})
        t1.append({"Catalyst": name, "Element": "χ^2", "Value": sci(r.chi2)})
        s2.append({"Sample": name, "Circuit": r.circuit.string, "N": len(r.f), "Outliers": r.n_outliers,
                   "f range (Hz)": f"{r.f.min():.4g} – {r.f.max():.4g}", "χ^2": r.chi2,
                   "Reduced χ^2": r.red_chi2, "RMS residual (%)": r.rms_pct, "R^2": r.r2, "AICc": r.aicc,
                   "KK RMS (%)": 100 * kk.rms if kk else np.nan, "KK max (%)": 100 * kk.max if kk else np.nan,
                   "Fit/KK": r.rms_pct / (100 * kk.rms) if kk else np.nan, "Valid": r.valid,
                   "Flags": "; ".join(r.flags)})
        for d in r.derived:
            s4.append({"Sample": name, "Element": d["element"], "C_eff (F)": d.get("C_eff", np.nan),
                       "τ (s)": d["tau"], "f_c (Hz)": d["f_c"]})
        if r.outliers is not None:
            for fx, zx in zip(r.f[r.outliers], r.Z[r.outliers]):
                out.append({"Sample": name, "f (Hz)": fx, "Z' (Ω)": zx.real, "-Z'' (Ω)": -zx.imag})
    s1 = [{"Sample": n, "Equivalent circuit": r.circuit.string, "Reading": r.circuit.explain(),
           "Selection": results.selection.get(n, "")} for n, r in results.fits.items()]
    return {"table1": pd.DataFrame(t1), "comparison": comparison_table(results, labels), "S1": pd.DataFrame(s1),
            "S2": pd.DataFrame(s2), "S3": pd.DataFrame(s3), "S4": pd.DataFrame(s4), "outliers": pd.DataFrame(out)}


def comparison_table(results, labels=None):
    """Elements as rows, samples as columns ('value ± SD unit'); χ² in the last row."""
    rows, order = {}, []
    for name, r in results.fits.items():
        for nm, v, s, rel, u in zip(r.circuit.names, r.p, r.sd, r.rel, r.circuit.units):
            lab, un = element_label(nm, labels), UNIT.get(u, u)
            key = lab + (f" ({un})" if un else "")
            if key not in rows:
                rows[key] = {}; order.append(key)
            rows[key][name] = value_pm(v, s) + (" ^a" if not np.isfinite(rel) or rel > 100 else "")
    rows["χ^2"] = {n: sci(r.chi2) for n, r in results.fits.items()}
    order.append("χ^2")
    df = pd.DataFrame([{"Element": k, **{n: rows[k].get(n, "–") for n in results.fits}} for k in order])
    return df


# ---------------------------------------------------------------------------------------
_SUP = str.maketrans("0123456789-+an", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺ᵃⁿ")


def plain(x):
    """Markup -> Unicode text for spreadsheets (R_s -> Rs, 10^-4 -> 10⁻⁴)."""
    if not isinstance(x, str):
        return x
    x = re.sub(r"\^([\-+]?\w+)", lambda m: m.group(1).translate(_SUP), x)
    return re.sub(r"_(\w+)", r"\1", x).replace(" ᵃ", "ᵃ")


def to_excel(results, path, labels=None):
    """Write all tables, fitted points and smooth fit curves to one .xlsx workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    T = tables(results, labels)
    wb = Workbook()
    head, fill = Font(name="Arial", bold=True, color="FFFFFF", size=10), PatternFill("solid", fgColor="1F4E78")
    body, line = Font(name="Arial", size=10), Border(bottom=Side(style="thin", color="BFBFBF"))

    def sheet(title, df, formulas=None):
        ws = wb.create_sheet(title)
        df = df.rename(columns=plain).map(plain)
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(1, j, c); cell.font, cell.fill = head, fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                v = None if isinstance(v, float) and not np.isfinite(v) else v
                c = ws.cell(i, j, v.item() if hasattr(v, "item") else v); c.font, c.border = body, line
            for j, fx in (formulas or {}).items():
                c = ws.cell(i, j, fx.format(r=i)); c.font, c.border = body, line
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[L(j)].width = min(max(len(str(c)), *(len(str(v)) for v in df[c]), 8) + 2, 48)
        ws.freeze_panes = "A2"
        return ws

    ws = wb.active; ws.title = "Read me"
    notes = [f"EIS equivalent-circuit fitting – {results.title or ', '.join(results.fits)}",
             "Fitting: complex non-linear least squares (modulus weighting), analytic Jacobian, multi-start; "
             "validity: Lin-KK (Schönleber 2014).",
             "Notation: Rs = R0; CPE-T = Q (S·sⁿ), CPE-P = n; Ws/Wo-R (Ω), Ws/Wo-T (s) = finite Warburg.",
             "ᵃ = not determined (relative standard error > 100 %).",
             "Robust fitting: points far from the trend were down-weighted (sheet 'Outliers').",
             "Sheet 'S3 95% CI': relative error, confidence limits and 'Determined' are Excel formulas."]
    for i, t in enumerate(notes, 1):
        ws.cell(i, 1, t).font = Font(name="Arial", bold=i == 1, size=12 if i == 1 else 10)
    ws.column_dimensions["A"].width = 110
    sheet("Comparison", T["comparison"])
    t1 = T["table1"].copy()
    t1["Catalyst"] = t1["Catalyst"].where(t1["Catalyst"] != t1["Catalyst"].shift(), "")
    sheet("Table 1", t1)
    sheet("S1 Circuits", T["S1"])
    sheet("S2 Goodness of fit", T["S2"])
    s3 = T["S3"].copy()
    sheet("S3 95% CI", s3, {6: "=E{r}/ABS(D{r})*100", 7: "=D{r}-1.96*E{r}", 8: "=D{r}+1.96*E{r}",
                            9: '=IF(F{r}>100,"no","yes")'})
    if len(T["S4"]):
        sheet("S4 Derived", T["S4"])
    pts, smooth = [], []
    for E, r in results.fits.items():
        pts.append(pd.DataFrame({"Sample": E, "f_Hz": r.f, "Zre_data": r.Z.real, "minusZim_data": -r.Z.imag,
                                 "Zre_fit": r.Zfit.real, "minusZim_fit": -r.Zfit.imag,
                                 "res_re_pct": 100 * r.res_re, "res_im_pct": 100 * r.res_im,
                                 "outlier": r.outliers if r.outliers is not None else False}))
        fd = np.logspace(np.log10(r.f.min()), np.log10(r.f.max()), 300); Zs = r.predict(fd)
        smooth.append(pd.DataFrame({"Sample": E, "f_Hz": fd, "Zre_fit": Zs.real, "minusZim_fit": -Zs.imag,
                                    "absZ_fit": np.abs(Zs), "minus_phase_deg": -np.angle(Zs, deg=True)}))
    sheet("Fitted points", pd.concat(pts))
    sheet("Smooth fit curves", pd.concat(smooth))
    if len(T["outliers"]):
        sheet("Outliers", T["outliers"])
    wb.save(path)
    return path


# ---------------------------------------------------------------------------------------
def to_docx(results, path, figures=None, labels=None):
    """Word report with Table 1 and Tables S1–S4 (booktabs style) and optional figures."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    def rich(par, text, bold=False, size=None, italic=False):
        for tok in re.split(r"(\^[−\-]?\w+|_[A-Za-z0-9]+)", str(text)):
            if not tok:
                continue
            if tok[0] == "^":
                run = par.add_run(tok[1:].replace("-", "−")); run.font.superscript = True
            elif tok[0] == "_":
                run = par.add_run(tok[1:]); run.font.subscript = True
            else:
                run = par.add_run(tok)
                if len(tok) == 1 and tok in "RCLW":
                    run.italic = True
            run.bold = bold; run.italic = run.italic or italic
            if size:
                run.font.size = Pt(size)

    def border(cell, **edges):
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders"); tcPr.append(b)
        for edge, sz in edges.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:color"), "000000")
            b.append(el)

    def table(doc, title, df, group_first=False, size=9, note=None, fmt=None):
        cap = doc.add_paragraph(); cap.paragraph_format.keep_with_next = True
        rich(cap, title, bold=True, size=11)
        cols = list(df.columns)
        t = doc.add_table(rows=1, cols=len(cols)); t.alignment = 1
        for j, c in enumerate(cols):
            p = t.rows[0].cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rich(p, c, size=size)
        prev = None
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for j, c in enumerate(cols):
                v = (fmt or {}).get(c, lambda x: f"{x:.4g}" if isinstance(x, float) else str(x))(row[c])
                if c == "Circuit":
                    v = v.replace("-", "‑")          # non-breaking hyphen keeps circuits on one line
                if group_first and j == 0:
                    raw, v = v, ("" if v == prev else v)
                    prev = raw
                p = cells[j].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                rich(p, v, size=size)
        for cell in t.rows[0].cells:
            border(cell, top=12, bottom=6)
        for cell in t.rows[-1].cells:
            border(cell, bottom=12)
        if note:
            rich(doc.add_paragraph(), note, size=8, italic=True)
        doc.add_paragraph()

    T = tables(results, labels)
    chem = lambda s: re.sub(r"(?<=[A-Za-z\)])(\d+|x)(?=[\(\s@]|$)", r"_\1", s)   # CoMoO4 -> CoMoO_4
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    rich(doc.add_paragraph(), "EIS fitting results" + (f" – {results.title}" if results.title else ""),
         bold=True, size=14)
    comp = T["comparison"].copy()
    comp.columns = [comp.columns[0]] + [chem(c) for c in comp.columns[1:]]
    nd = comp.iloc[:, 1:].apply(lambda c: c.str.contains(r"\^a")).any().any()
    table(doc, "Table 1. Equivalent-Circuit Parameters of the Samples", comp, size=9,
          note="Values are fitted value ± standard error. CPE-T and CPE-P are the CPE coefficient Q and exponent n "
               "(Z = 1/[Q(jω)^n]); W-R and W-T are the resistance and time constant of finite Warburg elements; "
               "χ^2 is the modulus-weighted chi-square." + (" ^a Not determined: relative standard error > 100 %."
                                                           if nd else ""))
    t1 = T["table1"].copy()
    t1["Catalyst"] = t1["Catalyst"].map(chem)
    table(doc, "Table 2. EIS Parameters of the Samples", t1, group_first=True,
          note="Values are fitted value ± standard error. CPE-T and CPE-P are the CPE coefficient Q and exponent n "
               "(Z = 1/[Q(jω)^n]); Ws/Wo-R and Ws/Wo-T are the resistance and time constant of the finite Warburg "
               "element; χ^2 is the modulus-weighted chi-square." +
               (" ^a Not determined: relative standard error > 100 %." if nd else ""))
    table(doc, "Table S1. Equivalent Circuits", T["S1"].drop(columns=["Reading"]),
          note="Circuit notation: '-' = series, p(A,B) = parallel.")
    sec = doc.add_section(); sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    s2 = T["S2"].drop(columns=["Flags", "Valid"])
    table(doc, "Table S2. Goodness of Fit and Kramers–Kronig Validation", s2, size=8,
          fmt={"χ^2": sci, "Reduced χ^2": sci, "RMS residual (%)": lambda x: f"{x:.3f}", "R^2": lambda x: f"{x:.4f}",
               "AICc": lambda x: f"{x:.1f}", "KK RMS (%)": lambda x: f"{x:.2f}", "KK max (%)": lambda x: f"{x:.2f}",
               "Fit/KK": lambda x: f"{x:.2f}", "Outliers": lambda x: str(int(x)), "N": lambda x: str(int(x))},
          note="RMS residual relative to |Z| (outliers excluded); Outliers = points down-weighted by the robust fit; "
               "KK = linear Kramers–Kronig test; Fit/KK ≈ 1: fit at noise level.")
    sec = doc.add_section(); sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    s3 = T["S3"].drop(columns=["Determined"])
    table(doc, "Table S3. Fitted Parameters with 95 % Confidence Intervals", s3, group_first=True, size=8,
          fmt={"Rel. error (%)": lambda x: f"{x:.1f}"},
          note="95 % CI = value ± 1.96 SD (asymptotic, from the Jacobian covariance).")
    if len(T["S4"]):
        table(doc, "Table S4. Derived Quantities", T["S4"], group_first=True,
              fmt={"C_eff (F)": lambda x: "–" if not np.isfinite(x) else sci(x, 3),
                   "τ (s)": lambda x: sci(x, 3), "f_c (Hz)": lambda x: sci(x, 3)},
              note="C_eff = Q^(1/n) R^((1−n)/n) (Hsu–Mansfeld); f_c = 1/(2πτ).")
    for k, (cap, fig_path) in enumerate((figures or {}).items(), 1):
        doc.add_picture(fig_path, width=Inches(6.5))
        rich(doc.add_paragraph(), f"Figure S{k}. {cap}", size=9)
    doc.save(path)
    return path
